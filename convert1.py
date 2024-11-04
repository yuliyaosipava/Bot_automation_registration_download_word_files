import os
from docx import Document
from openpyxl import Workbook

# Функция для извлечения данных из таблицы docx
def extract_table(docx_path):
    document = Document(docx_path)
    data = []
    for table in document.tables:
        for row in table.rows:
            data.append([cell.text for cell in row.cells])
    return data

# Функция для удаления точек из чисел и преобразования строк в числа
def clean_data(cell):
    try:
        # Удаляем точки и преобразуем в число
        return float(cell.replace('.', '').replace(',', '.'))
    except ValueError:
        # Если не удалось преобразовать в число, возвращаем оригинальную строку
        return cell

# Функция для записи данных в xlsx
def write_to_xlsx(data, xlsx_path):
    workbook = Workbook()
    sheet = workbook.active
    for row_data in data:
        cleaned_row = [clean_data(cell) for cell in row_data]
        sheet.append(cleaned_row)
    workbook.save(xlsx_path)

# Основная функция для преобразования всех файлов
def convert_files(docx_dir, xlsx_dir):
    for filename in os.listdir(docx_dir):
        if filename.endswith(".docx"):
            docx_path = os.path.join(docx_dir, filename)
            xlsx_filename = filename.replace(".docx", ".xlsx")
            xlsx_path = os.path.join(xlsx_dir, xlsx_filename)

            data = extract_table(docx_path)
            write_to_xlsx(data, xlsx_path)

# Укажите путь к директориям с вашими файлами
docx_dir = 'c:\\Users\\user\\Documents\\doci\\docs\\'  # Путь к папке с docx файлами
xlsx_dir = 'c:\\Users\\user\\Documents\\doci\\xls1\\'  # Путь к папке для сохранения xlsx файлов

# Преобразование файлов
convert_files(docx_dir, xlsx_dir)
